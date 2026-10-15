"""
S3.4 — the accountant's export: the month's Profit & Loss as XLSX and PDF (docs/v2/SPEC.md S3.4).
Done when "export opens and its totals equal /owner/report/monthly": the workbook is read back with openpyxl and every total is
compared with the report; the PDF is built uncompressed in tests so its text can be read.
"""
import io
import re
from datetime import date, datetime, timedelta

import pytest
from openpyxl import load_workbook

import auth
import models
from tests.fixtures import golden_restaurant as gr
from time_utils import eat_date, utcnow

API = "/api/v1/owner"
SEPT = datetime(2026, 9, 25, 10, 0)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _h(user):
    return {"Authorization": f"Bearer {auth.create_access_token({'sub': user.email, 'ver': 0})}"}


@pytest.fixture
def sept(db_session):
    """§7 #13 in September 2026: 412,000 sales, 131,800 ingredients, 103,000 wages (shifts), 112,000 other costs."""
    g = gr.build_skeleton(db_session, now=SEPT)
    gr.seed_pnl_month(db_session, g)
    return g


def _export(client, g, **params):
    return client.get(f"{API}/money/pnl/export", params=params, headers=_h(g.owner))


def _sheet(response):
    wb = load_workbook(io.BytesIO(response.content))
    return wb, {row[0].value: [c.value for c in row] for row in wb["P&L"].iter_rows() if row[0].value is not None}


# ── the workbook ─────────────────────────────────────────────────────────────

def test_the_export_is_a_real_xlsx_with_a_safe_filename(client, sept):
    res = _export(client, sept, month="2026-09")
    assert res.status_code == 200 and res.headers["content-type"] == XLSX
    disposition = res.headers["content-disposition"]
    assert re.fullmatch(r'attachment; filename="[A-Za-z0-9._-]+\.xlsx"', disposition), disposition
    assert "2026-09" in disposition and "pnl" in disposition
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["P&L", "Entries", "About"]
    assert res.content[:2] == b"PK"                                              # a zip: what an .xlsx is


def test_the_totals_in_the_workbook_are_the_worked_example(client, sept):
    _wb, rows = _sheet(_export(client, sept, month="2026-09"))
    assert rows["Sales"][1] == 412_000.0 and rows["Ingredients"][1] == 131_800.0 and rows["Left after ingredients"][1] == 280_200.0
    assert rows["Wages"][1] == 103_000.0 and rows["Other costs"][1] == 112_000.0 and rows["What you kept"][1] == 65_200.0
    assert rows["Sales"][2] == 100.0 and rows["Ingredients"][2] == 32.0 and rows["Wages"][2] == 25.0 and rows["What you kept"][2] == 15.8
    assert rows["Rent & premises"][1] == 60_000.0 and rows["Electricity (KPLC)"][1] == 14_000.0
    assert all(isinstance(rows[k][1], (int, float)) for k in ("Sales", "Ingredients", "Wages", "Other costs", "What you kept"))      # numbers, not text


def test_every_total_equals_the_monthly_report(client, sept):
    """The point of S3.4: the export and /owner/report/monthly are the same figures."""
    report = client.get(f"{API}/report/monthly", params={"month": "2026-09"}, headers=_h(sept.owner)).json()
    pnl = next(s for s in report["sections"] if s["key"] == "pnl")
    _wb, rows = _sheet(_export(client, sept, month="2026-09"))
    by_key = {l["key"]: abs(l["kes"]) for l in pnl["lines"]}
    assert (rows["Sales"][1], rows["Ingredients"][1], rows["Left after ingredients"][1], rows["Wages"][1], rows["Other costs"][1]) == \
        (by_key["sales"], by_key["ingredients"], by_key["left_after_ingredients"], by_key["wages"], by_key["other_costs"])
    assert rows["What you kept"][1] == pnl["lines"][-1]["kes"] == 65_200.0
    for sec in pnl["sections"]:
        assert rows[sec["name"]][1] == sec["kes"]
        for line in sec["lines"]:
            assert rows[line["name"]][1] == line["kes"], line["name"]
    assert pnl["period"]["label"] == "September 2026" and pnl["period"]["start"] == "2026-09-01"


def test_the_current_month_export_equals_the_default_monthly_report(client, db_session):
    g = gr.build_skeleton(db_session)                                          # the month of the real clock
    gr.seed_pnl_month(db_session, g)
    report = client.get(f"{API}/report/monthly", headers=_h(g.owner)).json()
    pnl = next(s for s in report["sections"] if s["key"] == "pnl")
    _wb, rows = _sheet(_export(client, g))
    assert rows["Sales"][1] == pnl["lines"][0]["kes"] and rows["What you kept"][1] == pnl["lines"][-1]["kes"]


def test_the_entries_sheet_lists_what_was_entered_and_adds_up(client, sept, db_session):
    db_session.add(models.Expense(restaurant_id=sept.rid, category_id=db_session.query(models.ExpenseCategory).filter_by(restaurant_id=sept.rid, key="repairs").one().id,
                                  amount_cents=999_900, paid_on=date(2026, 9, 12), method="cash", supplier="Voided Ltd", note="entered twice",
                                  voided_at=SEPT, void_reason="entered twice"))
    db_session.commit()
    wb, rows = _sheet(_export(client, sept, month="2026-09"))
    sheet = wb["Entries"]
    header = [c.value for c in sheet[1]]
    assert header == ["Date", "Cost line", "Section", "Paid to", "Method", "KES", "Paid by", "Note", "Receipt"]
    body = [[c.value for c in r] for r in sheet.iter_rows(min_row=2) if r[0].value not in (None, "Total entered")]
    assert len(body) == 8 and [r[0] for r in body] == sorted(r[0] for r in body)                      # the 8 §7 #13 costs; dates in order
    assert all(isinstance(r[5], (int, float)) for r in body) and "Voided Ltd" not in str([r[3] for r in body])
    total = next(r for r in sheet.iter_rows(min_row=2) if r[0].value == "Total entered")
    assert total[5].value == 112_000.0 == rows["Other costs"][1]                                      # wages are from the shifts, so nothing else was entered


def test_text_that_looks_like_a_formula_is_stored_as_text(client, sept, db_session):
    cat = db_session.query(models.ExpenseCategory).filter_by(restaurant_id=sept.rid, key="repairs").one()
    db_session.add(models.Expense(restaurant_id=sept.rid, category_id=cat.id, amount_cents=10_000, paid_on=date(2026, 9, 13), method="cash",
                                  supplier='=HYPERLINK("http://evil.example","pay here")', note="+1+1", receipt_ref="@SUM(A1)"))
    db_session.commit()
    wb, _rows = _sheet(_export(client, sept, month="2026-09"))
    texts = [c for r in wb["Entries"].iter_rows(min_row=2) for c in r if isinstance(c.value, str) and c.value[:1] in "=+-@"]
    assert len(texts) == 3 and all(c.data_type == "s" for c in texts)


def test_the_about_sheet_says_how_it_was_made(client, sept):
    wb, _ = _sheet(_export(client, sept, month="2026-09"))
    text = " ".join(str(c.value) for r in wb["About"].iter_rows() for c in r if c.value)
    assert "September 2026" in text and "paid" in text.lower() and "shift" in text.lower()
    assert "Golden Grill" in text                                                                    # the restaurant
    assert "Illustrative" not in text


def test_an_empty_month_still_exports(client, sept):
    res = _export(client, sept, month="2026-03")
    assert res.status_code == 200
    _wb, rows = _sheet(res)
    assert rows["Sales"][1] == 0 and rows["What you kept"][1] == 0
    about = " ".join(str(c.value) for r in load_workbook(io.BytesIO(res.content))["About"].iter_rows() for c in r if c.value)
    assert "Nothing recorded" in about


def test_estimates_and_gaps_are_written_into_the_workbook(client, db_session):
    g = gr.build_skeleton(db_session, now=SEPT)
    burger = g.items["Burger"]
    gr.add_order(db_session, g, [(burger, 1, 60_000)], datetime(2026, 9, 10, 9, 0), channel=models.DeliveryChannel.GLOVO, payment=models.PaymentMethod.CARD)
    db_session.commit()
    wb, rows = _sheet(_export(client, g, month="2026-09"))
    assert rows["Delivery app commissions"][1] == 150.0 and "estimated" in str(rows["Delivery app commissions"][3]).lower()
    notes = " ".join(str(c.value) for r in wb["About"].iter_rows() for c in r if c.value)
    assert "estimated" in notes.lower() and "Enter your rent" in notes


# ── the monthly report's `month` (what makes the totals comparable for any month) ───────────────

def test_the_monthly_report_of_a_chosen_month_covers_that_whole_month(client, sept):
    body = client.get(f"{API}/report/monthly", params={"month": "2026-09"}, headers=_h(sept.owner)).json()
    assert body["month"] == "2026-09" and body["range"] == "September 2026"
    sections = {s["key"]: s for s in body["sections"]}
    assert sections["pnl"]["period"]["start"] == "2026-09-01" and sections["pnl"]["period"]["end"] == "2026-09-30"
    assert sections["sales"]["sales"] == sections["pnl"]["lines"][0]["kes"] == 412_000.0 and sections["sales"]["orders"] == 410
    assert [d["name"] for d in sections["top_dishes"]["dishes"]][0] == "House set meal"
    other = client.get(f"{API}/report/monthly", params={"month": "2026-03"}, headers=_h(sept.owner)).json()
    assert {s["key"]: s["state"] for s in other["sections"]}["sales"] == "no_data"
    assert "month" not in client.get(f"{API}/report/monthly", headers=_h(sept.owner)).json()


@pytest.mark.parametrize("period,month", [("monthly", "2026-13"), ("monthly", "2026-9"), ("monthly", "2999-01"), ("monthly", "1999-12"), ("weekly", "2026-09"),
                                          ("daily", "2026-09"), ("yearly", "2026-09")])
def test_a_month_is_only_for_the_monthly_report_and_must_be_a_real_past_month(client, sept, period, month):
    assert client.get(f"{API}/report/{period}", params={"month": month}, headers=_h(sept.owner)).status_code == 422


# ── parameters, roles, tenants ───────────────────────────────────────────────

@pytest.mark.parametrize("params", [{"month": "2026-13"}, {"month": "2026-9"}, {"month": "september"}, {"month": "1999-12"}, {"month": "2999-01"},
                                    {"month": "2026-09", "format": "csv"}, {"format": "docx"}])
def test_bad_parameters_are_422(client, sept, params):
    assert _export(client, sept, **params).status_code == 422


def test_the_default_month_is_the_current_one_and_the_default_format_is_xlsx(client, sept):
    res = _export(client, sept)
    assert res.status_code == 200 and res.headers["content-type"] == XLSX
    assert str(eat_date(utcnow()).year) in res.headers["content-disposition"]


@pytest.mark.parametrize("role,status", [(models.StaffRole.MANAGER, 200), (models.StaffRole.CONTROLLER, 200), (models.StaffRole.SUPERVISOR, 403),
                                         (models.StaffRole.STOCKKEEPER, 403), (models.StaffRole.KITCHEN, 403), (models.StaffRole.WAITER, 403)])
def test_the_export_is_for_the_owner_the_manager_and_the_accountant(client, db_session, role, status):
    t = models.Tenant(name=f"T{role.value}")
    db_session.add(t)
    db_session.flush()
    db_session.add(models.Restaurant(tenant_id=t.id, name="R", address="x"))
    u = models.User(tenant_id=t.id, email=f"{role.value}@x.test", hashed_password=auth.get_password_hash("x"), role=models.Role.STAFF, staff_role=role, token_version=0)
    db_session.add(u)
    db_session.commit()
    for fmt in ("xlsx", "pdf"):
        assert client.get(f"{API}/money/pnl/export", params={"format": fmt}, headers=_h(u)).status_code == status


def test_anonymous_is_401_and_another_tenant_gets_only_its_own_books(client, sept, db_session):
    assert client.get(f"{API}/money/pnl/export").status_code == 401
    t = models.Tenant(name="Other")
    db_session.add(t)
    db_session.flush()
    db_session.add(models.Restaurant(tenant_id=t.id, name="Other Grill", address="x"))
    u = models.User(tenant_id=t.id, email="o@x.test", hashed_password=auth.get_password_hash("x"), role=models.Role.STAFF, staff_role=models.StaffRole.OWNER, token_version=0)
    db_session.add(u)
    db_session.commit()
    res = client.get(f"{API}/money/pnl/export", params={"month": "2026-09"}, headers=_h(u))
    wb, rows = _sheet(res)
    assert rows["Sales"][1] == 0 and rows["Other costs"][1] == 0
    blob = " ".join(str(c.value) for ws in wb for r in ws.iter_rows() for c in r if c.value)
    assert "Golden" not in blob and "Rent & premises" not in blob and "60000" not in blob and "Other Grill" in blob


# ── the PDF ──────────────────────────────────────────────────────────────────

def test_the_pdf_is_a_pdf_with_the_same_totals(client, sept, db_session):
    from owner.money.export import build_pdf, month_block
    res = _export(client, sept, month="2026-09", format="pdf")
    assert res.status_code == 200 and res.headers["content-type"] == "application/pdf" and res.content[:5] == b"%PDF-"
    assert re.fullmatch(r'attachment; filename="[A-Za-z0-9._-]+\.pdf"', res.headers["content-disposition"])
    from owner.states import OwnerContext
    user = db_session.get(models.User, sept.owner.id)
    ctx = OwnerContext(db_session, user, sept.restaurant)
    block, entries = month_block(ctx, date(2026, 9, 1))
    raw = build_pdf(block, entries, sept.restaurant.name, SEPT, compress=False).decode("latin-1")
    for expected in ("KES 412,000", "KES 131,800", "KES 280,200", "KES 103,000", "KES 112,000", "KES 65,200", "September 2026", "Out of every KES 100"):
        assert expected in raw, expected
    assert raw.count("/Type /Page\n") + raw.count("/Type /Page ") >= 1


def test_pdf_text_from_the_books_cannot_break_the_document(client, sept, db_session):
    cat = db_session.query(models.ExpenseCategory).filter_by(restaurant_id=sept.rid, key="repairs").one()
    db_session.add(models.Expense(restaurant_id=sept.rid, category_id=cat.id, amount_cents=10_000, paid_on=date(2026, 9, 13), method="cash",
                                  supplier="<b>Acme & <script>alert(1)</script> Ltd</b> — ünïcödé ≥ 5", note="</para><para>"))
    db_session.commit()
    res = _export(client, sept, month="2026-09", format="pdf")
    assert res.status_code == 200 and res.content[:5] == b"%PDF-"
    from owner.money.export import build_pdf, month_block
    from owner.states import OwnerContext
    block, entries = month_block(OwnerContext(db_session, db_session.get(models.User, sept.owner.id), sept.restaurant), date(2026, 9, 1))
    raw = build_pdf(block, entries, "Smith & Sons <Grill>", SEPT, compress=False).decode("latin-1")
    text = re.sub(r"\) Tj \(", "", raw)                       # the PDF draws a line in pieces; put the pieces back
    assert "<b>Acme & <script>alert\\(1\\)</script> Ltd</b>" in text and "</para><para>" in text      # printed as written, not read as markup
    assert "SMITH & SONS <GRILL>" in text
